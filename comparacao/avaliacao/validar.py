"""Confere o banco local sem indexação, embeddings ou chamadas de modelos."""
import csv
import hashlib
import io
import json
import re
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
EXPECTED = {
    "teste": {"factual": 25, "numerica": 25, "integracao": 20, "sintese": 15, "limites": 15},
    "desenvolvimento": {"factual": 5, "numerica": 5, "integracao": 4, "sintese": 3, "limites": 3},
}


def validate():
    from pypdf import PdfReader
    import openpyxl
    manifest = json.loads((HERE / "manifesto_corpus.json").read_text(encoding="utf-8"))
    corpus = ROOT / manifest["corpus"]
    snapshots = {item["arquivo"]: item for item in manifest["arquivos"]}
    actual_files = {p.name for p in corpus.iterdir() if p.is_file()}
    assert actual_files == set(snapshots), "Conjunto de arquivos do corpus mudou"
    for name, info in snapshots.items():
        assert hashlib.sha256((corpus / name).read_bytes()).hexdigest() == info["sha256"], f"Corpus mudou: {name}"
    readers, questions, ids, families = {}, set(), set(), {}
    source_count = 0
    reports = {}
    for split, expected in EXPECTED.items():
        records = [json.loads(line) for line in (HERE / f"{split}.jsonl").read_text(encoding="utf-8").splitlines()]
        assert Counter(item["categoria"] for item in records) == expected, f"Distribuição incorreta: {split}"
        reports[split] = dict(Counter(item["categoria"] for item in records))
        for record in records:
            assert record["id"] not in ids, "ID repetido"
            ids.add(record["id"])
            question = " ".join(record["pergunta"].lower().split())
            assert question not in questions, "Pergunta repetida"
            questions.add(question)
            assert record["schema_version"] == 1
            assert record["resposta_esperada"] and record["pontos_obrigatorios"] and record["fontes"]
            assert record["revisao_humana"]["status"] in {"pendente", "aprovada", "rejeitada"}
            assert record["comportamento_esperado"] in {
                "responder", "informacao_insuficiente", "pedir_esclarecimento", "corrigir_premissa"}
            families.setdefault(record["familia"], set()).add(split)
            for source in record["fontes"]:
                source_count += 1
                name = source["arquivo"]
                assert name in snapshots, f"Fonte fora do corpus: {name}"
                assert source["sha256"] == snapshots[name]["sha256"]
                path = corpus / name
                kind = source["tipo"]
                if name not in readers:
                    if kind == "pdf":
                        readers[name] = PdfReader(path)
                    elif kind == "xlsx":
                        readers[name] = openpyxl.load_workbook(path, read_only=True, data_only=True)
                    elif kind == "csv":
                        raw = path.read_bytes()
                        enc = "utf-16" if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8-sig"
                        try:
                            text = raw.decode(enc)
                        except UnicodeDecodeError:
                            text = raw.decode("cp1252")
                        dialect = csv.Sniffer().sniff(text[:10000], delimiters=";,\t")
                        readers[name] = list(csv.reader(io.StringIO(text), dialect))
                    else:
                        raise AssertionError(f"Tipo desconhecido: {kind}")
                reader = readers[name]
                if kind == "pdf":
                    page = source["pagina_pdf"]
                    assert 1 <= page <= len(reader.pages)
                    text = " ".join((reader.pages[page-1].extract_text() or "").split())
                    assert source["trecho"] in text, f"Trecho não encontrado: {name}, página {page}"
                elif kind == "csv":
                    line = source["linha"]
                    assert 2 <= line <= len(reader)
                    assert source["campos"] == dict(zip(reader[0], reader[line-1])), f"Linha divergente: {name}:{line}"
                else:
                    sheet = reader[source["aba"]]
                    assert str(sheet[source["celula"]].value or "") == source["valor_literal"], f"Célula divergente: {name}:{source['celula']}"
                    row = int(re.search(r"\d+$", source["celula"])[0])
                    assert str(sheet.cell(row,1).value or "") == source["rotulo_linha"]
                    for header in source["cabecalhos"]:
                        actual = [str(v) if v is not None else "" for v in next(sheet.iter_rows(
                            min_row=header["linha"],max_row=header["linha"],values_only=True))]
                        assert actual == header["valores"], f"Cabeçalho divergente: {name}:{source['aba']}"
    for reader in readers.values():
        close = getattr(reader,"close",None)
        if close: close()
    assert not any(len(parts)>1 for parts in families.values()), "Família repetida entre conjuntos"
    return {"status": "integridade_documental_verificada", "perguntas": len(questions),
            "referencias_conferidas": source_count, "arquivos_corpus": len(snapshots),
            "distribuicao": reports, "revisao_semantica_humana": "pendente"}


if __name__ == "__main__":
    print(json.dumps(validate(), ensure_ascii=False, indent=2))
