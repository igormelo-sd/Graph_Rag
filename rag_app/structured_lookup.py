"""Busca literal no corpus tabular completo, sem corte de linhas ou LLM."""
from __future__ import annotations
import csv
import hashlib
import io
import json
import re
import threading
from collections import defaultdict
from pathlib import Path

from domain_ontology import fold, facets, query_constraints
from llama_index.core.schema import TextNode, NodeWithScore

_STOP = set('qual quais foi foram entre frente primeiro trimestre ano codigo codigos ibge cnae classe cadastro territorial associado relacione seu sua uma para com dos das nos nas que como de do da no na ao em e o a os as pib valor valores corrente correntes'.split())
_CODE = re.compile(r'\b\d{2}\.?\d{2}-\d\b|\b\d{7}\b')

def terms(text):
    text=fold(text)
    text=re.sub(r'produto interno bruto','pib',text)
    return {t for t in re.findall(r'[a-z0-9]+',text) if len(t)>=3 and t not in _STOP}

def requested(question):
    return bool(re.search(r'cnae|ibge|municip|cadastro|classifica|planilha|\bcelula\b|\bpib\b|produto interno bruto',fold(question)))

class StructuredLookup:
    def __init__(self, directory):
        self.directory=Path(directory).resolve()
        self._cache={}
        self._lock=threading.Lock()

    def _read(self,path):
        records=[]
        source=path.relative_to(self.directory).as_posix()
        if path.suffix.lower()=='.csv':
            raw=path.read_bytes()
            if raw.startswith((b'\xff\xfe',b'\xfe\xff')): text=raw.decode('utf-16')
            else:
                try: text=raw.decode('utf-8-sig')
                except UnicodeDecodeError: text=raw.decode('cp1252')
            try: sep=csv.Sniffer().sniff(text[:65536],delimiters=';,\t').delimiter
            except csv.Error: sep=';' if ';' in text.splitlines()[0] else ','
            reader=csv.reader(io.StringIO(text),delimiter=sep)
            headers=next(reader,[])
            for number,row in enumerate(reader,2):
                if not any(row): continue
                if len(row)!=len(headers): raise ValueError(f'CSV inconsistente: {source}, registro {number}')
                records.append((headers,row,number,None,[]))
        else:
            import pandas as pd
            from ingestion import _spreadsheet_frame
            if path.suffix.lower()=='.xlsx':
                from openpyxl import load_workbook
                book=load_workbook(path,data_only=True)
                try:
                    sheets=[]
                    for sheet in book:
                        values=[list(row) for row in sheet.iter_rows(values_only=True)]
                        for area in sheet.merged_cells.ranges:
                            value=sheet.cell(area.min_row,area.min_col).value
                            for row in range(area.min_row-1,area.max_row):
                                for col in range(area.min_col-1,area.max_col): values[row][col]=value
                        sheets.append((sheet.title,pd.DataFrame(values),sheet))
                    for name,raw,sheet in sheets:
                        kept=[i for i in raw.columns if raw[i].notna().any()]
                        frame,_=_spreadsheet_frame(raw)
                        for index,row in frame.iterrows():
                            cells=[sheet.cell(int(index)+1,int(i)+1) for i in kept]
                            vals=[]
                            for value,cell in zip(row,cells):
                                if isinstance(value,(float,int)) and '%' in cell.number_format:
                                    decimals=len(re.search(r'\.(0+)',cell.number_format).group(1)) if re.search(r'\.(0+)',cell.number_format) else 0
                                    value=f'{value*100:.{decimals}f}%'.replace('.',',')
                                vals.append(str(value) if value is not None else '')
                            records.append((list(frame.columns),vals,int(index)+1,name,[c.coordinate for c in cells]))
                finally: book.close()
            else:
                with pd.ExcelFile(path) as book:
                    for name in book.sheet_names:
                        raw=book.parse(name,header=None)
                        kept=[i for i in raw.columns if raw[i].notna().any()]
                        frame,_=_spreadsheet_frame(raw)
                        from openpyxl.utils import get_column_letter
                        for index,row in frame.iterrows():
                            records.append((list(frame.columns),[str(v) for v in row],int(index)+1,name,[f'{get_column_letter(int(i)+1)}{int(index)+1}' for i in kept]))
        postings=defaultdict(set)
        cooked=[]
        sheet_notes=defaultdict(list)
        for headers,row,number,sheet,cells in records:
            value=next((str(v).strip() for v in row if str(v).strip()),'')
            if sheet and re.match(r'fonte\s*:|nota\b|dados\b|os dados\b|\(\d+\)',fold(value)):
                sheet_notes[sheet].append(f'{sheet}!linha {number}: '+ ' | '.join(str(v) for v in row if str(v).strip()))
        for headers,row,number,sheet,cells in records:
            body='\n'.join(f'{h}: {v}' for h,v in zip(headers,row) if str(v).strip())
            tokens=terms(' '.join(map(str,row)))
            if sheet: tokens |= terms(' '.join(headers))
            codes={re.sub(r'\D','',str(v)) for h,v in zip(headers,row) if re.search(r'cod|codigo|cnae',fold(h)) and re.fullmatch(r'[\d.\-]+',str(v).strip())}
            tokens |= codes
            names={fold(str(v)).strip() for h,v in zip(headers,row) if re.search(r'municipio|localidade',fold(h)) and str(v).strip()}
            years={str(v) for h,v in zip(headers,row) if re.search(r'\bano\b|periodo|trimestres',fold(h)) and re.fullmatch(r'(19|20)\d{2}',str(v))}
            quarters={int(str(v)[0]) for h,v in zip(headers,row) if 'trimestres' in fold(h) and re.fullmatch(r'[1-4][ºo°]',str(v))}
            quarters |= { {'I':1,'II':2,'III':3,'IV':4}[str(v)] for h,v in zip(headers,row) if 'trimestres' in fold(h) and str(v) in {'I','II','III','IV'} }
            if sheet and years:
                body+='\n[Notas literais da mesma aba]\n'+'\n'.join(sheet_notes[sheet])
            header_years=set(re.findall(r'\b(?:19|20)\d{2}\b',' '.join(headers)))
            cooked.append({'text':body,'row':number,'sheet':sheet,'cells':cells,'tokens':tokens,'codes':codes,'names':names,'years':years or header_years,'quarters':quarters,'headers':headers})
            for token in tokens: postings[token].add(len(cooked)-1)
        return cooked,postings,any(x['years'] for x in cooked)

    def retrieve(self,question,limit=8):
        if not requested(question) or not self.directory.is_dir(): return []
        query=terms(question)
        codes={re.sub(r'\D','',v) for v in _CODE.findall(question)}
        query |= codes
        years=set(re.findall(r'\b(?:19|20)\d{2}\b',question))
        raw_question=fold(question)
        quarter=re.search(r'\b(primeiro|segundo|terceiro|quarto|[1-4])[ºo°]?\s*(?:trimestre|t\s*20\d{2})',raw_question)
        quarter_number=None
        if quarter:
            quarter_number=int(quarter[1]) if quarter[1].isdigit() else {'primeiro':1,'segundo':2,'terceiro':3,'quarto':4}[quarter[1]]
        wanted_region=set(query_constraints(question)['dimensions']['region'])
        addresses=set(re.findall(r'\b[A-Z]{1,3}[1-9]\d{0,6}\b',question)) if re.search(r'celula|célula|!',question,re.I) else set()
        catalogue=bool(re.search(r'fontes|arquivos|bases',fold(question)) and re.search(r'classifica|cadastro|territori',fold(question)))
        selected=[]
        # Um snapshot por consulta; mudanças invalidam somente a fonte afetada.
        with self._lock:
            paths=sorted(p for p in self.directory.rglob('*') if p.suffix.lower() in {'.csv','.xlsx','.xls'} and p.is_file() and not p.name.startswith('~$'))
            for path in paths:
                stat=path.stat(); key=(stat.st_mtime_ns,stat.st_size)
                cached=self._cache.get(path)
                if cached is None or cached[0]!=key:
                    cached=(key,self._read(path)); self._cache[path]=cached
                records,postings,temporal=cached[1]
                named_sheets={r['sheet'] for r in records if r['sheet'] and re.search(r'(?<!\w)'+re.escape(fold(r['sheet']))+r'(?!\w)',fold(question))}
                if catalogue and records:
                    headers=records[0]['headers']
                    if re.search(r'cnae|secao|classe|municipio|ibge|localidade',fold(' '.join(headers))):
                        body='Arquivo: '+path.relative_to(self.directory).as_posix()+'\nCabeçalhos literais: '+'; '.join(headers)
                        if len(records)<=20: body+='\n'+ '\n---\n'.join(r['text'] for r in records)
                        source=path.relative_to(self.directory).as_posix()
                        node=TextNode(id_='catalogue:'+hashlib.sha256((source+body).encode()).hexdigest()[:24],text=body,metadata={'type':'table','source_file':source,'page':'cabeçalho','retrieval_method':'structured_catalogue'})
                        selected.append((20,NodeWithScore(node=node,score=.9)))
                    continue
                candidates=set().union(*(postings.get(t,set()) for t in query)) if query else set()
                if addresses: candidates={i for i,r in enumerate(records) if addresses.intersection(r['cells'])}
                ranked=[]
                for index in candidates:
                    r=records[index]
                    if named_sheets and r['sheet'] not in named_sheets: continue
                    if r['sheet'] and fold(r['sheet']).startswith('indice') and not named_sheets: continue
                    if addresses and not addresses.intersection(r['cells']): continue
                    if codes and not codes.intersection(r['codes']): continue
                    if years and temporal and not years.intersection(r['years']): continue
                    if quarter_number and r['quarters'] and quarter_number not in r['quarters']: continue
                    if quarter_number and r['sheet'] and not r['quarters']: continue
                    got_region=facets(' '.join(r['headers']))['region'] if r['sheet'] else set()
                    if wanted_region and got_region and wanted_region.isdisjoint(got_region): continue
                    if wanted_region=={'estado_sp'} and r['sheet'] and 'estado de sao paulo' not in fold(' '.join(r['headers'])): continue
                    matched=query & r['tokens']
                    # Ano sozinho não identifica uma linha.
                    if not (matched-years-codes) and not codes and not addresses: continue
                    name=any(re.search(r'(?<!\w)'+re.escape(n)+r'(?!\w)',fold(question)) for n in r['names'])
                    score=len(matched)+100*bool(codes or addresses)+40*name
                    if 'pib' in fold(question) and re.search(r'produto interno bruto|\bpib\b',fold(r['text'])): score+=15
                    if r['sheet']:
                        header=fold(' '.join(r['headers']))
                        if re.search(r'variacao|variou|crescimento|cresceu|taxa',raw_question) and 'taxas de crescimento' in header: score+=20
                        if 'trimestre anterior' in raw_question and 'trimestre imediatamente anterior' in header: score+=30
                        if 'ajuste sazonal' in raw_question and 'ajuste sazonal' in header: score+=20
                    if 'per capita' not in fold(question) and 'per capita' in fold(r['text']): score-=20
                    ranked.append((score,index,bool(codes or addresses or name)))
                for score,index,exact in sorted(ranked,key=lambda x:(-x[0],x[1]))[:3]:
                    r=records[index]; source=path.relative_to(self.directory).as_posix()
                    location=f"{r['sheet']}!linha {r['row']}" if r['sheet'] else f"linha {r['row']}"
                    nodeid='structured:'+hashlib.sha256(f"{source}:{key}:{location}:{r['text']}".encode()).hexdigest()[:24]
                    metadata={'type':'table','source_file':source,'page':location,'row':r['row'],'sheet':r['sheet'] or '', 'cells':json.dumps(r['cells']), 'retrieval_method':'structured_exact' if exact else 'structured_lexical','table_total_rows':len(records),'table_truncated':False}
                    selected.append((score,NodeWithScore(node=TextNode(id_=nodeid,text=r['text'],metadata=metadata),score=1.0 if exact else .75)))
            live=set(paths)
            self._cache={p:v for p,v in self._cache.items() if p in live}
        return [n for _,n in sorted(selected,key=lambda x:-x[0])[:limit]]
